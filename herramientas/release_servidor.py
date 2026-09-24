"""Construye y verifica releases reproducibles del servidor Edge.

La lista permitida contiene solamente código. Estado operativo, secretos y
artefactos de desarrollo se excluyen recursivamente.
"""

from __future__ import annotations

import argparse
import base64
import binascii
import csv
from dataclasses import dataclass
import datetime as dt
from email import policy as email_policy
from email.parser import BytesParser
import hashlib
import io
import json
import os
from pathlib import Path, PurePosixPath
import re
import stat
import subprocess
import sys
import tempfile
from urllib.parse import unquote, urlparse
from urllib.request import url2pathname
import uuid
import warnings
import zipfile


PRODUCTO = "LosTocayosPOS-Servidor"
FORMATO_MANIFIESTO = 2
POLITICA_CONTENIDO = "edge-server-v2"
ENTRADA_MANIFIESTO = "_release/manifest.json"
DESTINO_RELEASE = {
    "implementation": "cp",
    "python": "3.13",
    "abi": "cp313",
    "platform": "win_amd64",
    "bits": 64,
}
TAMANO_MAXIMO_ARCHIVO = 256 * 1024 * 1024
TAMANO_MAXIMO_PAYLOAD = 1024 * 1024 * 1024
TAMANO_MAXIMO_MANIFIESTO = 16 * 1024 * 1024
TAMANO_MAXIMO_XLSX_DESCOMPRIMIDO = 64 * 1024 * 1024
ENTRADAS_MAXIMAS_XLSX = 1024
TAMANO_MAXIMO_WHEEL_DESCOMPRIMIDO = 512 * 1024 * 1024
ENTRADAS_MAXIMAS_WHEEL = 4096
RUTA_CATALOGO_HISTORICO = "datos/Listado-Productos.xlsx"
ENCABEZADOS_CATALOGO_HISTORICO = (
    "Codigo", "Nombre", "Categoria", "Precio", "Destino impresion",
)

RUTAS_REQUERIDAS = (
    ".env.example", "VERSION", "requirements.txt", "requirements-lock.txt",
    "manage.py", "servicio_windows.py", "instalar-servicio-lan.ps1",
    "instalar-servidor.ps1", "actualizar-servidor.ps1",
    "actualizar-laboratorio-desde-release.ps1",
    "aprovisionar-sucursal.ps1", "reparar-permisos-servidor.ps1",
    "iniciar-servicio-lan.ps1", "verificar-servicio-lan.ps1",
    "respaldar-db-sqlite.ps1", "certs/prod-ca-2021.crt",
    RUTA_CATALOGO_HISTORICO,
    "catalogo", "contracts", "herramientas", "impresion",
    "personas", "pos", "tests", "ventas",
)
ARCHIVOS_CONTRATO_REQUERIDOS = (
    ".env.example", "VERSION", "requirements.txt", "requirements-lock.txt",
    "manage.py", "servicio_windows.py", "instalar-servicio-lan.ps1",
    "instalar-servidor.ps1", "actualizar-servidor.ps1",
    "actualizar-laboratorio-desde-release.ps1",
    "aprovisionar-sucursal.ps1", "reparar-permisos-servidor.ps1",
    "iniciar-servicio-lan.ps1", "verificar-servicio-lan.ps1",
    "respaldar-db-sqlite.ps1", "certs/prod-ca-2021.crt",
    RUTA_CATALOGO_HISTORICO,
    "herramientas/validar_despliegue.py", "pos/settings.py",
    "personas/identidad.py",
    "personas/management/commands/aprovisionar_sucursal.py",
    "personas/management/commands/inicializar_operacion_sucursal.py",
    "personas/management/commands/verificar_identidad_local.py",
)
PREFIJOS_CONTRATO_REQUERIDOS = (
    "catalogo/", "contracts/", "herramientas/", "impresion/", "personas/", "pos/",
    "tests/", "ventas/",
)
RUTAS_OPCIONALES = (
    "README.md", "DESPLIEGUE_WINDOWS.md", "iniciar-local.ps1",
)
ARCHIVOS_EN_DIRECTORIOS_RESERVADOS = frozenset({RUTA_CATALOGO_HISTORICO})
DIRECTORIOS_EXCLUIDOS = {
    ".cache", ".git", ".hg", ".mypy_cache", ".nox", ".pytest_cache",
    ".ruff_cache", ".svn", ".tox", ".venv", "__pycache__", "backups",
    "cache", "data", "datos", "desktop", "htmlcov", "logs",
    "media", "node_modules", "release", "releases", "runtime", "secrets",
    "staticfiles", "temp", "tmp",
}
ARCHIVOS_SECRETOS = {
    ".coverage", ".env", "credentials.json", "id_dsa", "id_ecdsa",
    "id_ed25519", "id_rsa", "secrets.json",
}
SUFIJOS_EXCLUIDOS = (
    ".bak", ".cred", ".credentials", ".db", ".jks", ".key", ".log",
    ".p12", ".pem", ".pfx", ".pyc", ".pyo", ".secret", ".sqlite",
    ".sqlite-journal", ".sqlite-shm", ".sqlite-wal", ".sqlite3",
    ".sqlite3-journal", ".sqlite3-shm", ".sqlite3-wal", ".tmp",
)
PATRON_VERSION = re.compile(r"^[0-9A-Za-z][0-9A-Za-z._+-]{0,63}$")
PATRON_COMMIT = re.compile(r"^(?:[0-9a-fA-F]{40}|[0-9a-fA-F]{64})$")
PATRON_SHA256 = re.compile(r"^[0-9a-f]{64}$")


class ErrorRelease(RuntimeError):
    """La release no puede construirse o no supera una verificación."""


@dataclass(frozen=True)
class ArtefactosRelease:
    zip: Path
    manifiesto: Path
    sumas: Path
    sha256_zip: str
    sha256_manifiesto: str


@dataclass(frozen=True)
class ArchivoPayload:
    ruta: str
    origen: Path
    tamano: int
    sha256: str


def _es_directorio_excluido(nombre: str) -> bool:
    nombre = nombre.casefold()
    return (
        nombre in DIRECTORIOS_EXCLUIDOS
        or nombre.startswith(".venv-")
        or nombre.endswith(".egg-info")
    )


def _es_archivo_excluido(nombre: str) -> bool:
    nombre = nombre.casefold()
    if nombre == ".env.example":
        return False
    return (
        nombre in ARCHIVOS_SECRETOS
        or nombre.startswith(".env.")
        or nombre.endswith(SUFIJOS_EXCLUIDOS)
    )


def _validar_ruta_relativa(valor: str) -> str:
    normalizada = valor.replace("\\", "/")
    ruta = PurePosixPath(normalizada)
    reservados = {"con", "prn", "aux", "nul", "clock$"}
    if (
        not normalizada
        or normalizada != ruta.as_posix()
        or ruta.is_absolute()
        or normalizada.startswith("/")
        or any(parte in {"", ".", ".."} for parte in ruta.parts)
        or any(
            any(ord(caracter) < 32 or caracter in '<>:"|?*' for caracter in parte)
            or len(parte.encode("utf-16-le")) // 2 > 255
            or parte.endswith((" ", "."))
            for parte in ruta.parts
        )
        or any(
            parte.split(".", 1)[0].casefold()
            in reservados | {"conin$", "conout$"}
            or re.fullmatch(
                r"(?:com|lpt)[1-9¹²³]",
                parte.split(".", 1)[0],
                re.IGNORECASE,
            )
            for parte in ruta.parts
        )
    ):
        raise ErrorRelease(f"Ruta relativa no segura: {valor!r}.")
    return ruta.as_posix()


def _es_enlace(ruta: Path) -> bool:
    if ruta.is_symlink():
        return True
    es_union = getattr(ruta, "is_junction", None)
    return bool(es_union and es_union())


def _rechazar_enlaces_en_componentes(ruta: Path, descripcion: str) -> None:
    """Rechaza enlaces o junctions tanto en la ruta como en sus padres."""

    absoluta = Path(os.path.abspath(ruta))
    for componente in (absoluta, *absoluta.parents):
        if _es_enlace(componente):
            raise ErrorRelease(
                f"{descripcion} no puede atravesar enlaces simbólicos o uniones."
            )


def _resolver_dentro_del_origen(origen: Path, ruta: Path, descripcion: str) -> Path:
    """Resuelve una entrada sólo si toda su ruta es real y permanece en origen."""

    _rechazar_enlaces_en_componentes(ruta, descripcion)
    try:
        resuelta = ruta.resolve(strict=True)
    except (OSError, RuntimeError) as exc:
        raise ErrorRelease(f"No se pudo resolver {descripcion.lower()}.") from exc
    try:
        resuelta.relative_to(origen)
    except ValueError as exc:
        raise ErrorRelease(f"{descripcion} escapa del directorio de origen.") from exc
    return resuelta


def _validar_catalogo_historico(origen: Path | io.BytesIO, descripcion: str) -> None:
    """Comprueba que el XLSX contiene el contrato posicional usado al importar."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ErrorRelease(
            "Se requiere openpyxl para validar el catálogo histórico."
        ) from exc

    libro = None
    try:
        if isinstance(origen, io.BytesIO):
            origen.seek(0)
        with zipfile.ZipFile(origen, "r") as paquete_opc:
            infos = paquete_opc.infolist()
            nombres = {info.filename for info in infos}
            if (
                not infos
                or len(infos) > ENTRADAS_MAXIMAS_XLSX
                or len(nombres) != len(infos)
                or "[Content_Types].xml" not in nombres
                or "xl/workbook.xml" not in nombres
                or any(info.flag_bits & 0x1 for info in infos)
                or any(
                    info.file_size > TAMANO_MAXIMO_ARCHIVO
                    for info in infos
                )
                or sum(info.file_size for info in infos)
                > TAMANO_MAXIMO_XLSX_DESCOMPRIMIDO
                or paquete_opc.testzip() is not None
            ):
                raise ErrorRelease(
                    f"{descripcion} no es un paquete XLSX/OPC íntegro y acotado."
                )
        if isinstance(origen, io.BytesIO):
            origen.seek(0)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            libro = load_workbook(origen, read_only=True, data_only=True)
        if "Productos" not in libro.sheetnames:
            raise ErrorRelease(
                f"{descripcion} no contiene la hoja obligatoria Productos."
            )
        hoja = libro["Productos"]
        encabezados = tuple(
            next(hoja.iter_rows(min_row=1, max_row=1, max_col=5, values_only=True))
        )
        if encabezados != ENCABEZADOS_CATALOGO_HISTORICO:
            raise ErrorRelease(
                f"{descripcion} no contiene las columnas esperadas: "
                + ", ".join(ENCABEZADOS_CATALOGO_HISTORICO)
                + "."
            )
        if not any(
            fila[1] is not None and str(fila[1]).strip()
            for fila in hoja.iter_rows(min_row=2, max_col=5, values_only=True)
        ):
            raise ErrorRelease(
                f"{descripcion} no contiene productos importables."
            )
    except ErrorRelease:
        raise
    except Exception as exc:
        raise ErrorRelease(
            f"{descripcion} no es un archivo XLSX/OPC utilizable."
        ) from exc
    finally:
        if libro is not None:
            libro.close()


def _validar_wheel(ruta: Path, descripcion: str) -> None:
    """Valida estructura, metadatos y límites de un wheel sin confiar en su nombre."""

    try:
        from packaging.utils import canonicalize_name, parse_wheel_filename
        from packaging.version import Version

        from packaging.tags import parse_tag

        nombre_archivo, version_archivo, _, tags_archivo = parse_wheel_filename(ruta.name)
        with zipfile.ZipFile(ruta, "r") as wheel:
            infos = wheel.infolist()
            nombres = [info.filename for info in infos]
            if (
                not infos
                or len(infos) > ENTRADAS_MAXIMAS_WHEEL
                or len({nombre.casefold() for nombre in nombres}) != len(nombres)
                or any(info.flag_bits & 0x1 for info in infos)
                or any(
                    info.file_size > TAMANO_MAXIMO_ARCHIVO
                    or info.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
                    for info in infos
                )
                or sum(info.file_size for info in infos)
                > TAMANO_MAXIMO_WHEEL_DESCOMPRIMIDO
            ):
                raise ErrorRelease(f"{descripcion} no es un wheel íntegro y acotado.")

            for info in infos:
                nombre = info.filename[:-1] if info.is_dir() else info.filename
                _validar_ruta_relativa(nombre)
                tipo = (info.external_attr >> 16) & 0o170000
                permitidos = {0, stat.S_IFDIR if info.is_dir() else stat.S_IFREG}
                if tipo not in permitidos:
                    raise ErrorRelease(f"{descripcion} contiene una entrada no regular.")
                leidos = 0
                if not info.is_dir():
                    with wheel.open(info, "r") as contenido:
                        for bloque in iter(lambda: contenido.read(1024 * 1024), b""):
                            leidos += len(bloque)
                            if leidos > info.file_size or leidos > TAMANO_MAXIMO_ARCHIVO:
                                raise ErrorRelease(f"{descripcion} excede sus límites declarados.")
                if leidos != info.file_size:
                    raise ErrorRelease(f"{descripcion} contiene tamaños incoherentes.")

            candidatos = [
                nombre.removesuffix("/METADATA")
                for nombre in nombres
                if nombre.endswith(".dist-info/METADATA")
            ]
            if len(candidatos) != 1:
                raise ErrorRelease(f"{descripcion} no contiene un único dist-info.")
            prefijo = candidatos[0]
            try:
                base_dist_info = prefijo.removesuffix(".dist-info")
                nombre_dist_info, version_dist_info = base_dist_info.rsplit("-", 1)
                dist_info_valido = (
                    len(PurePosixPath(prefijo).parts) == 1
                    and canonicalize_name(nombre_dist_info)
                    == canonicalize_name(nombre_archivo)
                    and Version(version_dist_info) == version_archivo
                )
            except ValueError:
                dist_info_valido = False
            if not dist_info_valido:
                raise ErrorRelease(
                    f"{descripcion} contiene un dist-info que no coincide con su nombre."
                )
            for nombre in nombres:
                partes = PurePosixPath(
                    nombre[:-1] if nombre.endswith("/") else nombre
                ).parts
                indices_dist_info = [
                    indice
                    for indice, parte in enumerate(partes)
                    if parte.endswith(".dist-info")
                ]
                if indices_dist_info and (
                    indices_dist_info != [0] or partes[0] != prefijo
                ):
                    raise ErrorRelease(
                        f"{descripcion} contiene más de un árbol dist-info."
                    )
            requeridos = {
                f"{prefijo}/METADATA", f"{prefijo}/WHEEL", f"{prefijo}/RECORD"
            }
            if not requeridos.issubset(nombres):
                raise ErrorRelease(f"{descripcion} no contiene los metadatos obligatorios.")
            mensaje = BytesParser(policy=email_policy.default).parsebytes(
                wheel.read(f"{prefijo}/METADATA")
            )
            versiones_metadata = mensaje.get_all("Metadata-Version", [])
            nombres_metadata = mensaje.get_all("Name", [])
            versiones_paquete = mensaje.get_all("Version", [])
            if (
                mensaje.defects
                or len(versiones_metadata) != 1
                or len(nombres_metadata) != 1
                or len(versiones_paquete) != 1
                or not re.fullmatch(r"[1-9]\d*\.\d+", str(versiones_metadata[0]))
                or canonicalize_name(str(nombres_metadata[0]))
                != canonicalize_name(nombre_archivo)
                or Version(str(versiones_paquete[0])) != version_archivo
            ):
                raise ErrorRelease(f"{descripcion} no coincide con sus metadatos.")

            mensaje_wheel = BytesParser(policy=email_policy.default).parsebytes(
                wheel.read(f"{prefijo}/WHEEL")
            )
            versiones_wheel = mensaje_wheel.get_all("Wheel-Version", [])
            generadores_wheel = mensaje_wheel.get_all("Generator", [])
            ubicaciones_wheel = mensaje_wheel.get_all("Root-Is-Purelib", [])
            wheel_version = (
                str(versiones_wheel[0]) if len(versiones_wheel) == 1 else ""
            )
            root_is_purelib = (
                str(ubicaciones_wheel[0]).casefold()
                if len(ubicaciones_wheel) == 1
                else ""
            )
            if (
                mensaje_wheel.defects
                or len(versiones_wheel) != 1
                or len(generadores_wheel) != 1
                or not str(generadores_wheel[0]).strip()
                or len(ubicaciones_wheel) != 1
                or not re.fullmatch(r"1\.\d+", wheel_version)
                or root_is_purelib not in {"true", "false"}
            ):
                raise ErrorRelease(f"{descripcion} contiene metadatos WHEEL inválidos.")
            tags_metadata = set()
            for valor_tag in mensaje_wheel.get_all("Tag", []):
                tags_metadata.update(parse_tag(valor_tag))
            if not tags_metadata or tags_metadata != set(tags_archivo):
                raise ErrorRelease(
                    f"{descripcion} no coincide con las etiquetas de plataforma de WHEEL."
                )

            ruta_record = f"{prefijo}/RECORD"
            try:
                texto_record = wheel.read(ruta_record).decode("utf-8")
                filas_record = list(csv.reader(io.StringIO(texto_record), strict=True))
            except (UnicodeDecodeError, csv.Error) as exc:
                raise ErrorRelease(f"{descripcion} contiene un RECORD ilegible.") from exc
            registros: dict[str, tuple[str, str, str]] = {}
            vistas_record: set[str] = set()
            for fila in filas_record:
                if len(fila) != 3:
                    raise ErrorRelease(f"{descripcion} contiene una fila RECORD inválida.")
                nombre_record = _validar_ruta_relativa(fila[0])
                clave_record = nombre_record.casefold()
                if clave_record in vistas_record:
                    raise ErrorRelease(f"{descripcion} repite una ruta en RECORD.")
                vistas_record.add(clave_record)
                registros[nombre_record] = (fila[0], fila[1], fila[2])

            archivos_wheel = {nombre for nombre in nombres if not nombre.endswith("/")}
            firmas_record = {f"{prefijo}/RECORD.jws", f"{prefijo}/RECORD.p7s"}
            esperados_record = archivos_wheel - firmas_record
            if set(registros) != esperados_record:
                raise ErrorRelease(
                    f"{descripcion} no declara exactamente sus archivos en RECORD."
                )
            for nombre_record, (_, hash_record, tamano_record) in registros.items():
                if nombre_record == ruta_record:
                    if hash_record or tamano_record:
                        raise ErrorRelease(
                            f"{descripcion} debe dejar vacíos hash y tamaño de RECORD."
                        )
                    continue
                if not hash_record.startswith("sha256=") or not tamano_record.isascii() or not tamano_record.isdecimal():
                    raise ErrorRelease(
                        f"{descripcion} debe usar SHA-256 y tamaños decimales en RECORD."
                    )
                contenido = wheel.read(nombre_record)
                huella = hash_record.partition("=")[2]
                try:
                    decodificada = base64.urlsafe_b64decode(
                        huella + "=" * (-len(huella) % 4)
                    )
                except (ValueError, binascii.Error) as exc:
                    raise ErrorRelease(
                        f"{descripcion} contiene un hash RECORD inválido."
                    ) from exc
                huella_canonica = base64.urlsafe_b64encode(
                    hashlib.sha256(contenido).digest()
                ).rstrip(b"=").decode("ascii")
                if (
                    decodificada != hashlib.sha256(contenido).digest()
                    or huella != huella_canonica
                    or int(tamano_record) != len(contenido)
                ):
                    raise ErrorRelease(
                        f"{descripcion} no coincide con los hashes o tamaños de RECORD."
                    )
    except ErrorRelease:
        raise
    except (OSError, ValueError, RuntimeError, UnicodeError, zipfile.BadZipFile) as exc:
        raise ErrorRelease(f"{descripcion} está dañado o no es un wheel válido.") from exc


def _entorno_pip_seguro() -> dict[str, str]:
    entorno = {
        nombre: valor
        for nombre, valor in os.environ.items()
        if not nombre.upper().startswith(("PYTHON", "PIP_"))
        and nombre.upper() not in {"VIRTUAL_ENV", "__PYVENV_LAUNCHER__"}
    }
    entorno.update(
        {
            "PYTHONNOUSERSITE": "1",
            "PYTHONDONTWRITEBYTECODE": "1",
            "PYTHONUTF8": "1",
            "PIP_CONFIG_FILE": os.devnull,
            "PIP_DISABLE_PIP_VERSION_CHECK": "1",
            "PIP_NO_INPUT": "1",
        }
    )
    return entorno


def _validar_lock_exacto(requirements_path: Path) -> list[str]:
    """Devuelve los pins activos para Windows CPython 3.13, sin fuentes externas."""

    from packaging.markers import default_environment
    from packaging.requirements import Requirement
    from packaging.utils import canonicalize_name
    from packaging.version import Version

    if not requirements_path.is_file() or requirements_path.stat().st_size > 1024 * 1024:
        raise ErrorRelease("requirements-lock.txt no es un archivo regular y acotado.")
    entorno_destino = default_environment()
    entorno_destino.update(
        {
            "implementation_name": "cpython",
            "implementation_version": "3.13.0",
            "os_name": "nt",
            "platform_machine": "AMD64",
            "platform_python_implementation": "CPython",
            "platform_release": "",
            "platform_system": "Windows",
            "platform_version": "",
            "python_full_version": "3.13.0",
            "python_version": "3.13",
            "sys_platform": "win32",
            "extra": "",
        }
    )
    vistos: set[str] = set()
    activos: list[str] = []
    activos_por_nombre: set[str] = set()
    try:
        lineas = requirements_path.read_text(encoding="utf-8-sig").splitlines()
    except (OSError, UnicodeError) as exc:
        raise ErrorRelease("requirements-lock.txt no es texto UTF-8 legible.") from exc
    for numero, cruda in enumerate(lineas, start=1):
        linea = cruda.strip()
        if not linea or linea.startswith("#"):
            continue
        try:
            requisito = Requirement(linea)
        except ValueError as exc:
            raise ErrorRelease(
                f"Pin inválido en requirements-lock.txt, línea {numero}."
            ) from exc
        especificadores = list(requisito.specifier)
        if (
            requisito.url
            or requisito.extras
            or len(especificadores) != 1
            or especificadores[0].operator != "=="
            or especificadores[0].version.endswith(".*")
        ):
            raise ErrorRelease(
                "requirements-lock.txt sólo admite un pin name==version por línea, "
                "sin extras, rangos, comodines, URLs ni opciones de pip."
            )
        nombre = canonicalize_name(requisito.name)
        if nombre in vistos:
            raise ErrorRelease(f"requirements-lock.txt repite el paquete {nombre}.")
        vistos.add(nombre)
        try:
            version = str(Version(especificadores[0].version))
            activo = requisito.marker is None or requisito.marker.evaluate(entorno_destino)
        except (KeyError, ValueError) as exc:
            raise ErrorRelease(
                f"Marcador o versión inválida para {requisito.name}."
            ) from exc
        if activo:
            activos.append(f"{requisito.name}=={version}")
            activos_por_nombre.add(nombre)
    if not activos or "pywin32" not in activos_por_nombre:
        raise ErrorRelease(
            "El lock debe activar pywin32 y al menos un paquete para Windows CPython 3.13."
        )
    return activos


def _ruta_wheel_desde_url(url: object, carpeta: Path) -> Path:
    if not isinstance(url, str):
        raise ErrorRelease("El reporte de pip contiene una URL no textual.")
    analizada = urlparse(url)
    if (
        analizada.scheme.casefold() != "file"
        or analizada.hostname not in {None, "", "localhost"}
        or analizada.username is not None
        or analizada.password is not None
        or analizada.query
        or analizada.fragment
    ):
        raise ErrorRelease("pip seleccionó una dependencia fuera del wheelhouse local.")
    try:
        ruta_texto = url2pathname(unquote(analizada.path))
        candidata = Path(ruta_texto).resolve(strict=True)
    except (OSError, RuntimeError, ValueError) as exc:
        raise ErrorRelease("pip seleccionó una ruta local no resoluble.") from exc
    carpeta_real = carpeta.resolve(strict=True)
    if (
        candidata.parent != carpeta_real
        or not candidata.is_file()
        or not candidata.name.casefold().endswith(".whl")
    ):
        raise ErrorRelease("pip seleccionó un archivo que no pertenece al wheelhouse plano.")
    return candidata


def _validar_wheelhouse_exacto(carpeta: Path, requirements_path: Path) -> None:
    """Exige pins canónicos y wheels exactos para Windows CPython 3.13 x64."""

    from packaging.requirements import Requirement
    from packaging.utils import canonicalize_name
    from packaging.version import Version

    pins_activos = _validar_lock_exacto(requirements_path)
    pines_esperados: dict[str, Version] = {}
    for pin in pins_activos:
        requisito = Requirement(pin)
        especificador = next(iter(requisito.specifier))
        pines_esperados[canonicalize_name(requisito.name)] = Version(
            especificador.version
        )
    with tempfile.TemporaryDirectory(prefix="tocayos-pip-resolve-") as temporal:
        raiz_temporal = Path(temporal)
        reporte = raiz_temporal / "report.json"
        lock_destino = raiz_temporal / "requirements-win-cp313.txt"
        lock_destino.write_text("\n".join(pins_activos) + "\n", encoding="utf-8")
        resultado = subprocess.run(
            [
                sys.executable, "-I", "-m", "pip", "install", "--dry-run",
                "--ignore-installed", "--disable-pip-version-check", "--no-index",
                "--only-binary=:all:", "--find-links", str(carpeta),
                "--platform", DESTINO_RELEASE["platform"],
                "--implementation", DESTINO_RELEASE["implementation"],
                "--python-version", DESTINO_RELEASE["python"],
                "--abi", DESTINO_RELEASE["abi"],
                "--report", str(reporte), "-r", str(lock_destino),
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=_entorno_pip_seguro(),
            check=False,
        )
        if resultado.returncode:
            raise ErrorRelease(
                "El wheelhouse no resuelve completamente requirements-lock.txt "
                "para Windows CPython 3.13 x64."
            )
        try:
            contenido = json.loads(reporte.read_text(encoding="utf-8"))
            instalaciones = contenido["install"]
        except (OSError, UnicodeError, json.JSONDecodeError, KeyError, TypeError) as exc:
            raise ErrorRelease("pip no produjo un reporte de resolución válido.") from exc
        seleccionados: set[Path] = set()
        paquetes_seleccionados: dict[str, Version] = {}
        for instalacion in instalaciones:
            try:
                url = instalacion["download_info"]["url"]
                metadatos = instalacion["metadata"]
                nombre_paquete = canonicalize_name(metadatos["name"])
                version_paquete = Version(metadatos["version"])
            except (KeyError, TypeError, ValueError) as exc:
                raise ErrorRelease("El reporte de pip contiene una fuente no reconocida.") from exc
            ruta_seleccionada = _ruta_wheel_desde_url(url, carpeta)
            if ruta_seleccionada in seleccionados:
                raise ErrorRelease("pip seleccionó dos veces el mismo wheel.")
            if nombre_paquete in paquetes_seleccionados:
                raise ErrorRelease("pip seleccionó dos veces el mismo paquete.")
            seleccionados.add(ruta_seleccionada)
            paquetes_seleccionados[nombre_paquete] = version_paquete
        if paquetes_seleccionados != pines_esperados:
            raise ErrorRelease(
                "Cada paquete resuelto debe tener un pin exacto y activo en "
                "requirements-lock.txt."
            )
        entradas = list(carpeta.iterdir())
        if any(
            not entrada.is_file() or not entrada.name.casefold().endswith(".whl")
            for entrada in entradas
        ):
            raise ErrorRelease("El wheelhouse debe ser plano y contener sólo wheels.")
        entregados = {entrada.resolve(strict=True) for entrada in entradas}
        if seleccionados != entregados:
            raise ErrorRelease(
                "El wheelhouse debe contener exactamente los wheels seleccionados por "
                "requirements-lock.txt, sin paquetes adicionales."
            )


def _validar_contrato_payload(rutas: set[str]) -> None:
    normalizadas = {ruta.casefold() for ruta in rutas}
    faltantes = [
        ruta for ruta in ARCHIVOS_CONTRATO_REQUERIDOS
        if ruta.casefold() not in normalizadas
    ]
    faltantes.extend(
        prefijo + "*"
        for prefijo in PREFIJOS_CONTRATO_REQUERIDOS
        if not any(ruta.startswith(prefijo.casefold()) for ruta in normalizadas)
    )
    if faltantes:
        raise ErrorRelease(
            "La release no cumple el contrato mínimo; faltan: "
            + ", ".join(faltantes)
        )


def _sha256_archivo(ruta: Path) -> str:
    digest = hashlib.sha256()
    with ruta.open("rb") as archivo:
        for bloque in iter(lambda: archivo.read(1024 * 1024), b""):
            digest.update(bloque)
    return digest.hexdigest()


def _agregar_archivo(
    encontrados: dict[str, ArchivoPayload],
    origen: Path,
    relativa: str,
    total_actual: int,
) -> int:
    relativa = _validar_ruta_relativa(relativa)
    if _es_archivo_excluido(PurePosixPath(relativa).name):
        return total_actual
    if _es_enlace(origen):
        raise ErrorRelease(f"No se permiten enlaces simbólicos o uniones: {relativa}.")
    if not origen.is_file():
        raise ErrorRelease(f"La entrada no es un archivo regular: {relativa}.")
    tamano = origen.stat().st_size
    if tamano > TAMANO_MAXIMO_ARCHIVO:
        raise ErrorRelease(f"El archivo excede el límite permitido: {relativa}.")
    total = total_actual + tamano
    if total > TAMANO_MAXIMO_PAYLOAD:
        raise ErrorRelease("El contenido de la release excede el límite permitido.")
    clave = relativa.casefold()
    if clave in encontrados:
        if encontrados[clave].origen != origen:
            raise ErrorRelease(f"Dos archivos colisionan por mayúsculas/minúsculas: {relativa}.")
        return total_actual
    encontrados[clave] = ArchivoPayload(
        ruta=relativa,
        origen=origen,
        tamano=tamano,
        sha256=_sha256_archivo(origen),
    )
    return total


def _recopilar_payload(origen: Path, rutas: tuple[str, ...]) -> list[ArchivoPayload]:
    encontrados: dict[str, ArchivoPayload] = {}
    total = 0
    for incluida in rutas:
        relativa_raiz = _validar_ruta_relativa(incluida)
        partes = PurePosixPath(relativa_raiz).parts
        if (
            any(_es_directorio_excluido(parte) for parte in partes)
            and relativa_raiz not in ARCHIVOS_EN_DIRECTORIOS_RESERVADOS
        ):
            raise ErrorRelease(f"No se puede incluir una ruta reservada: {incluida}.")
        entrada_solicitada = origen.joinpath(*partes)
        _rechazar_enlaces_en_componentes(
            entrada_solicitada, f"La ruta incluida {incluida}"
        )
        if not entrada_solicitada.exists():
            raise ErrorRelease(f"Falta la ruta que debe incluirse: {incluida}.")
        entrada = _resolver_dentro_del_origen(
            origen, entrada_solicitada, f"La ruta incluida {incluida}"
        )
        if entrada.is_file():
            if _es_archivo_excluido(entrada.name):
                raise ErrorRelease(f"No se puede incluir un archivo sensible: {incluida}.")
            total = _agregar_archivo(encontrados, entrada, relativa_raiz, total)
            continue
        if not entrada.is_dir():
            raise ErrorRelease(f"La ruta incluida no es archivo ni directorio: {incluida}.")

        for carpeta, subdirectorios, archivos in os.walk(entrada, followlinks=False):
            carpeta_path = _resolver_dentro_del_origen(
                origen, Path(carpeta), "Una carpeta del payload"
            )
            conservados: list[str] = []
            for nombre in sorted(subdirectorios, key=str.casefold):
                candidata = carpeta_path / nombre
                if _es_directorio_excluido(nombre):
                    continue
                _resolver_dentro_del_origen(
                    origen, candidata, "Una carpeta del payload"
                )
                conservados.append(nombre)
            subdirectorios[:] = conservados
            for nombre in sorted(archivos, key=str.casefold):
                if _es_archivo_excluido(nombre):
                    continue
                archivo = _resolver_dentro_del_origen(
                    origen, carpeta_path / nombre, "Un archivo del payload"
                )
                relativa = archivo.relative_to(origen).as_posix()
                total = _agregar_archivo(encontrados, archivo, relativa, total)
    return sorted(encontrados.values(), key=lambda item: item.ruta.encode("utf-8"))


def _agregar_wheelhouse(
    payload: list[ArchivoPayload],
    wheelhouse: Path | None,
    requirements_path: Path,
) -> list[ArchivoPayload]:
    if wheelhouse is None:
        return payload
    solicitada = Path(wheelhouse)
    _rechazar_enlaces_en_componentes(solicitada, "El wheelhouse")
    carpeta = solicitada.resolve(strict=True)
    if not carpeta.is_dir():
        raise ErrorRelease("El wheelhouse debe ser un directorio real.")

    encontrados = {item.ruta.casefold(): item for item in payload}
    total = sum(item.tamano for item in payload)
    paquetes = sorted(carpeta.iterdir(), key=lambda item: item.name.casefold())
    if not paquetes:
        raise ErrorRelease("El wheelhouse está vacío.")
    for paquete in paquetes:
        if _es_enlace(paquete) or not paquete.is_file():
            raise ErrorRelease("El wheelhouse debe ser plano y contener sólo archivos regulares.")
        if not paquete.name.casefold().endswith(".whl"):
            raise ErrorRelease(f"El wheelhouse sólo admite wheels (.whl): {paquete.name}.")
        if not zipfile.is_zipfile(paquete):
            raise ErrorRelease(f"El archivo no es un wheel ZIP válido: {paquete.name}.")
        _validar_wheel(paquete, f"El wheel {paquete.name}")
        total = _agregar_archivo(
            encontrados,
            paquete,
            f"wheelhouse/{paquete.name}",
            total,
        )
    _validar_wheelhouse_exacto(carpeta, requirements_path)
    return sorted(encontrados.values(), key=lambda item: item.ruta.encode("utf-8"))


def _entorno_git_sin_reemplazos() -> dict[str, str]:
    entorno = os.environ.copy()
    # Los objetos replace son locales y no forman parte del commit publicado.
    entorno["GIT_NO_REPLACE_OBJECTS"] = "1"
    return entorno


def _ejecutar_git(origen: Path, argumentos: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        ["git", "-C", str(origen), *argumentos],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=_entorno_git_sin_reemplazos(),
        check=False,
    )


def _rechazar_payload_ignorado_git(
    origen: Path,
    payload: list[ArchivoPayload],
) -> None:
    """Impide que archivos locales ignorados se filtren dentro de rutas permitidas."""

    if _ejecutar_git(origen, ["rev-parse", "--verify", "HEAD"]).returncode:
        return
    resultado = subprocess.run(
        ["git", "-C", str(origen), "check-ignore", "-z", "--stdin"],
        input="".join(item.ruta + "\0" for item in payload),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=_entorno_git_sin_reemplazos(),
        check=False,
    )
    if resultado.returncode not in {0, 1}:
        raise ErrorRelease("No se pudo comprobar la política Git del payload.")
    ignoradas = [linea for linea in resultado.stdout.split("\0") if linea]
    if ignoradas:
        raise ErrorRelease(
            "La lista permitida contiene archivos ignorados por Git; "
            "muévelos fuera del código o inclúyelos deliberadamente en el repositorio: "
            + ", ".join(ignoradas)
        )


def _fragmentar_rutas_git(
    payload: list[ArchivoPayload],
    limite_caracteres: int = 6000,
) -> list[list[str]]:
    """Mantiene las invocaciones bajo el límite de línea de comandos de Windows."""

    fragmentos: list[list[str]] = []
    actual: list[str] = []
    longitud = 0
    for item in payload:
        adicional = len(item.ruta) + 3
        if actual and longitud + adicional > limite_caracteres:
            fragmentos.append(actual)
            actual = []
            longitud = 0
        actual.append(item.ruta)
        longitud += adicional
    if actual:
        fragmentos.append(actual)
    return fragmentos


def _rechazar_transformaciones_git(
    origen: Path,
    payload: list[ArchivoPayload],
) -> None:
    """Rechaza filtros capaces de ocultar bytes distintos detrás de un estado clean."""

    for rutas in _fragmentar_rutas_git(payload):
        resultado = subprocess.run(
            [
                "git", "-C", str(origen), "check-attr", "-z", "--stdin",
                "filter", "ident", "working-tree-encoding",
            ],
            input="".join(ruta + "\0" for ruta in rutas),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=_entorno_git_sin_reemplazos(),
            check=False,
        )
        if resultado.returncode:
            raise ErrorRelease("No se pudieron comprobar los atributos Git del payload.")
        campos = resultado.stdout.split("\0")
        if campos and campos[-1] == "":
            campos.pop()
        if len(campos) % 3:
            raise ErrorRelease("Git devolvió atributos de rutas no reconocidos.")
        for indice in range(0, len(campos), 3):
            ruta, atributo, valor = campos[indice:indice + 3]
            if valor not in {"unspecified", "unset"}:
                raise ErrorRelease(
                    "El payload usa una transformación Git no permitida "
                    f"({atributo}={valor}) en {ruta}."
                )


def _verificar_payload_contra_head(
    origen: Path,
    payload: list[ArchivoPayload],
    commit: str,
) -> None:
    """Prueba que un payload declarado limpio corresponde realmente a HEAD."""

    head = _ejecutar_git(origen, ["rev-parse", "--verify", "HEAD"])
    if head.returncode or head.stdout.strip().lower() != commit:
        raise ErrorRelease("HEAD cambió durante la construcción de la release.")

    _rechazar_transformaciones_git(origen, payload)

    indice = _ejecutar_git(origen, ["ls-files", "-v", "-z", "--cached"])
    if indice.returncode:
        raise ErrorRelease("No se pudieron comprobar las marcas del índice Git.")
    marcas: dict[str, str] = {}
    for registro in indice.stdout.split("\0"):
        if not registro:
            continue
        if len(registro) < 3 or registro[1] != " ":
            raise ErrorRelease("Git devolvió una entrada de índice no reconocida.")
        marca, ruta = registro[0], registro[2:]
        if ruta in marcas:
            raise ErrorRelease(f"El índice Git contiene etapas ambiguas para {ruta}.")
        marcas[ruta] = marca

    arbol = _ejecutar_git(origen, ["ls-tree", "-r", "-z", "--full-tree", commit])
    if arbol.returncode:
        raise ErrorRelease("No se pudo leer el árbol Git de HEAD.")
    blobs: dict[str, str] = {}
    for registro in arbol.stdout.split("\0"):
        if not registro:
            continue
        try:
            metadatos, ruta = registro.split("\t", 1)
            _, tipo, oid = metadatos.split(" ", 2)
        except ValueError as exc:
            raise ErrorRelease("Git devolvió una entrada de árbol no reconocida.") from exc
        if tipo == "blob":
            blobs[ruta] = oid

    for item in payload:
        marca = marcas.get(item.ruta)
        if marca is None or item.ruta not in blobs:
            raise ErrorRelease(
                f"El payload limpio contiene un archivo no registrado en HEAD: {item.ruta}."
            )
        # H es el estado normal de una entrada cacheada. h (assume-unchanged),
        # S (skip-worktree) y cualquier estado especial invalidan la atestación.
        if marca != "H":
            raise ErrorRelease(
                "El índice Git contiene marcas assume-unchanged, skip-worktree "
                f"o un estado no permitido para {item.ruta}."
            )

    hashes_trabajo: dict[str, str] = {}
    for rutas in _fragmentar_rutas_git(payload):
        resultado = _ejecutar_git(origen, ["hash-object", "--", *rutas])
        hashes = resultado.stdout.splitlines()
        if resultado.returncode or len(hashes) != len(rutas):
            raise ErrorRelease("No se pudo comparar el payload con HEAD.")
        hashes_trabajo.update(zip(rutas, hashes, strict=True))

    diferentes = [
        item.ruta
        for item in payload
        if hashes_trabajo[item.ruta].lower() != blobs[item.ruta].lower()
    ]
    if diferentes:
        muestra = ", ".join(diferentes[:5])
        if len(diferentes) > 5:
            muestra += ", ..."
        raise ErrorRelease(
            "El payload declarado limpio no coincide con HEAD: " + muestra
        )


def _resolver_git(
    origen: Path,
    commit_solicitado: str | None,
    rutas: tuple[str, ...],
    permitir_sucio: bool,
) -> tuple[str, bool]:
    resultado_head = _ejecutar_git(origen, ["rev-parse", "--verify", "HEAD"])
    tiene_git = resultado_head.returncode == 0
    head = resultado_head.stdout.strip().lower() if tiene_git else ""
    if commit_solicitado:
        if not PATRON_COMMIT.fullmatch(commit_solicitado):
            raise ErrorRelease("El commit debe ser un hash Git completo de 40 o 64 caracteres.")
        commit = commit_solicitado.lower()
        if tiene_git and commit != head:
            raise ErrorRelease("El commit indicado no coincide con HEAD del árbol empaquetado.")
    else:
        if not tiene_git or not PATRON_COMMIT.fullmatch(head):
            raise ErrorRelease(
                "No se pudo resolver el commit. Usa un checkout Git o indica --commit."
            )
        commit = head

    # Un árbol exportado puede declarar a qué commit pretende corresponder, pero
    # sin metadatos Git no es posible comprobarlo ni afirmar que esté limpio.
    sucio = not tiene_git
    if tiene_git:
        estado = _ejecutar_git(
            origen, ["status", "--porcelain=v1", "--untracked-files=all", "--", *rutas]
        )
        if estado.returncode:
            raise ErrorRelease("No se pudo comprobar el estado Git del contenido.")
        sucio = bool(estado.stdout)
        if sucio and not permitir_sucio:
            raise ErrorRelease(
                "El contenido de la release tiene cambios sin commit; "
                "confírmalos o usa --allow-dirty de forma explícita."
            )
    return commit, sucio


def _resolver_epoch(origen: Path, commit: str, valor: int | None) -> int:
    if valor is None and os.environ.get("SOURCE_DATE_EPOCH"):
        try:
            valor = int(os.environ["SOURCE_DATE_EPOCH"])
        except ValueError as exc:
            raise ErrorRelease("SOURCE_DATE_EPOCH debe ser un entero.") from exc
    if valor is None:
        resultado = _ejecutar_git(origen, ["show", "-s", "--format=%ct", commit])
        if resultado.returncode == 0:
            try:
                valor = int(resultado.stdout.strip())
            except ValueError:
                valor = None
    if valor is None:
        raise ErrorRelease(
            "No se pudo fijar la fecha reproducible; indica --source-date-epoch."
        )
    if valor < 315532800 or valor > 4354819199:
        raise ErrorRelease("La fecha reproducible debe estar entre 1980 y 2107.")
    return valor


def _fecha_iso(epoch: int) -> str:
    return (
        dt.datetime.fromtimestamp(epoch, tz=dt.timezone.utc)
        .replace(microsecond=0)
        .isoformat()
        .replace("+00:00", "Z")
    )


def _fecha_zip(epoch: int) -> tuple[int, int, int, int, int, int]:
    fecha = dt.datetime.fromtimestamp(epoch, tz=dt.timezone.utc)
    return (
        fecha.year, fecha.month, fecha.day, fecha.hour, fecha.minute,
        fecha.second // 2 * 2,
    )


def _bytes_json(valor: object) -> bytes:
    return (
        json.dumps(valor, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")


def _info_zip(nombre: str, fecha: tuple[int, int, int, int, int, int]) -> zipfile.ZipInfo:
    info = zipfile.ZipInfo(filename=nombre, date_time=fecha)
    info.compress_type = zipfile.ZIP_DEFLATED
    info.create_system = 3
    permisos = 0o100755 if nombre.endswith(".sh") else 0o100644
    info.external_attr = permisos << 16
    info.extra = b""
    info.comment = b""
    return info


def _escribir_entrada(
    paquete: zipfile.ZipFile,
    item: ArchivoPayload,
    fecha: tuple[int, int, int, int, int, int],
) -> None:
    digest = hashlib.sha256()
    escritos = 0
    with item.origen.open("rb") as fuente, paquete.open(
        _info_zip(item.ruta, fecha), "w", force_zip64=True
    ) as destino:
        for bloque in iter(lambda: fuente.read(1024 * 1024), b""):
            digest.update(bloque)
            escritos += len(bloque)
            destino.write(bloque)
    if escritos != item.tamano or digest.hexdigest() != item.sha256:
        raise ErrorRelease(f"El archivo cambió durante el empaquetado: {item.ruta}.")


def _validar_version(version: str) -> str:
    if not PATRON_VERSION.fullmatch(version) or ".." in version:
        raise ErrorRelease(
            "La versión debe tener 1-64 caracteres alfanuméricos, '.', '_', '+', o '-'."
        )
    return version


def _version_autoritativa(origen: Path) -> str:
    ruta = origen / "VERSION"
    if not ruta.is_file() or _es_enlace(ruta):
        raise ErrorRelease("Falta el archivo regular VERSION autoritativo.")
    try:
        lineas = ruta.read_text(encoding="ascii").splitlines()
    except (OSError, UnicodeDecodeError) as exc:
        raise ErrorRelease("VERSION debe ser texto ASCII legible.") from exc
    if len(lineas) != 1 or lineas[0] != lineas[0].strip():
        raise ErrorRelease("VERSION debe contener exactamente una línea canónica.")
    return _validar_version(lineas[0])


def crear_release(
    *,
    origen: Path,
    destino: Path,
    version: str,
    commit: str | None = None,
    source_date_epoch: int | None = None,
    permitir_sucio: bool = False,
    sobrescribir: bool = False,
    rutas_incluidas: tuple[str, ...] | None = None,
    wheelhouse: Path | None = None,
) -> ArtefactosRelease:
    """Crea ZIP, manifiesto y sumas sin incluir estado local."""

    origen_solicitado = Path(origen)
    _rechazar_enlaces_en_componentes(origen_solicitado, "El origen")
    origen = origen_solicitado.resolve(strict=True)
    if not origen.is_dir():
        raise ErrorRelease("El origen debe ser un directorio real, no un enlace.")
    version = _validar_version(version)
    if version != _version_autoritativa(origen):
        raise ErrorRelease("La versión solicitada no coincide con VERSION.")
    if rutas_incluidas is None:
        faltantes = [ruta for ruta in RUTAS_REQUERIDAS if not (origen / ruta).exists()]
        if faltantes:
            raise ErrorRelease("Faltan rutas requeridas: " + ", ".join(faltantes))
        rutas = RUTAS_REQUERIDAS + tuple(
            ruta for ruta in RUTAS_OPCIONALES if (origen / ruta).exists()
        )
    else:
        rutas = tuple(rutas_incluidas)
        if not rutas:
            raise ErrorRelease("Debe incluirse al menos una ruta.")

    canonicas: list[str] = []
    vistas: set[str] = set()
    for ruta in rutas:
        canonica = _validar_ruta_relativa(ruta)
        if canonica.casefold() not in vistas:
            canonicas.append(canonica)
            vistas.add(canonica.casefold())
    rutas = tuple(canonicas)

    commit_resuelto, sucio = _resolver_git(origen, commit, rutas, permitir_sucio)
    epoch = _resolver_epoch(origen, commit_resuelto, source_date_epoch)
    _validar_lock_exacto(origen / "requirements-lock.txt")
    payload_fuente = _recopilar_payload(origen, rutas)
    _rechazar_payload_ignorado_git(origen, payload_fuente)
    payload = _agregar_wheelhouse(
        payload_fuente,
        wheelhouse,
        origen / "requirements-lock.txt",
    )
    if not payload:
        raise ErrorRelease("La lista permitida no produjo ningún archivo.")
    _validar_contrato_payload({item.ruta for item in payload})
    catalogo = next(
        (
            item for item in payload_fuente
            if item.ruta == RUTA_CATALOGO_HISTORICO
        ),
        None,
    )
    if catalogo is None:
        raise ErrorRelease(
            "El catálogo histórico debe usar su ruta y mayúsculas canónicas."
        )
    _validar_catalogo_historico(
        catalogo.origen, "El catálogo histórico incluido"
    )
    if not sucio:
        _verificar_payload_contra_head(
            origen, payload_fuente, commit_resuelto
        )

    nombre_base = f"{PRODUCTO}-{version}"
    destino_solicitado = Path(destino)
    _rechazar_enlaces_en_componentes(destino_solicitado, "El destino")
    destino = destino_solicitado.resolve()
    if destino.exists() and not destino.is_dir():
        raise ErrorRelease("El destino debe ser un directorio real.")
    destino.mkdir(parents=True, exist_ok=True)
    zip_final = destino / f"{nombre_base}.zip"
    manifiesto_final = destino / f"{nombre_base}.manifest.json"
    sumas_final = destino / f"{nombre_base}.sha256"
    finales = (zip_final, manifiesto_final, sumas_final)
    existentes = [ruta.name for ruta in finales if ruta.exists()]
    if existentes and not sobrescribir:
        raise ErrorRelease(
            "La release es inmutable y ya existen artefactos: " + ", ".join(existentes)
        )

    manifiesto = {
        "artifact_name": zip_final.name,
        "content_policy": POLITICA_CONTENIDO,
        "dependency_bundle": "wheelhouse" if wheelhouse is not None else "none",
        "created_utc": _fecha_iso(epoch),
        "file_count": len(payload),
        "files": [
            {"path": item.ruta, "sha256": item.sha256, "size": item.tamano}
            for item in payload
        ],
        "format_version": FORMATO_MANIFIESTO,
        "payload_size": sum(item.tamano for item in payload),
        "product": PRODUCTO,
        "release_version": version,
        "source_commit": commit_resuelto,
        "source_date_epoch": epoch,
        "source_dirty": sucio,
        "target": DESTINO_RELEASE.copy(),
    }
    bytes_manifiesto = _bytes_json(manifiesto)
    sufijo = f".tmp-{os.getpid()}-{uuid.uuid4().hex}"
    zip_temporal = destino / f".{zip_final.name}{sufijo}"
    manifiesto_temporal = destino / f".{manifiesto_final.name}{sufijo}"
    sumas_temporal = destino / f".{sumas_final.name}{sufijo}"
    temporales = (zip_temporal, manifiesto_temporal, sumas_temporal)
    try:
        fecha = _fecha_zip(epoch)
        with zipfile.ZipFile(
            zip_temporal, "x", compression=zipfile.ZIP_DEFLATED,
            compresslevel=9, allowZip64=True,
        ) as paquete:
            paquete.comment = b""
            for item in payload:
                _escribir_entrada(paquete, item, fecha)
            paquete.writestr(
                _info_zip(ENTRADA_MANIFIESTO, fecha),
                bytes_manifiesto,
                compress_type=zipfile.ZIP_DEFLATED,
                compresslevel=9,
            )
        # Segunda atestación: detecta cambios posteriores a la lectura de un
        # archivo y antes de publicar el artefacto (ventana TOCTOU).
        if not sucio:
            _verificar_payload_contra_head(
                origen, payload_fuente, commit_resuelto
            )
        manifiesto_temporal.write_bytes(bytes_manifiesto)
        sha_zip = _sha256_archivo(zip_temporal)
        sha_manifiesto = _sha256_archivo(manifiesto_temporal)
        sumas_temporal.write_bytes(
            (
                f"{sha_zip}  {zip_final.name}\n"
                f"{sha_manifiesto}  {manifiesto_final.name}\n"
            ).encode("ascii")
        )
        os.replace(zip_temporal, zip_final)
        os.replace(manifiesto_temporal, manifiesto_final)
        # Las sumas se publican al final: los otros artefactos ya están completos.
        os.replace(sumas_temporal, sumas_final)
        for ruta in finales:
            os.utime(ruta, (epoch, epoch))
        return ArtefactosRelease(
            zip_final, manifiesto_final, sumas_final, sha_zip, sha_manifiesto
        )
    finally:
        for temporal in temporales:
            try:
                temporal.unlink()
            except FileNotFoundError:
                pass


def _json_sin_duplicados(datos: bytes) -> object:
    def construir(pares: list[tuple[str, object]]) -> dict[str, object]:
        resultado: dict[str, object] = {}
        for clave, valor in pares:
            if clave in resultado:
                raise ErrorRelease(f"Clave JSON duplicada en manifiesto: {clave}.")
            resultado[clave] = valor
        return resultado

    try:
        return json.loads(datos.decode("utf-8"), object_pairs_hook=construir)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ErrorRelease("El manifiesto no es JSON UTF-8 válido.") from exc


def _leer_sumas(ruta: Path) -> dict[str, str]:
    try:
        lineas = ruta.read_text(encoding="ascii").splitlines()
    except (OSError, UnicodeDecodeError) as exc:
        raise ErrorRelease("No se pudo leer el archivo SHA-256.") from exc
    sumas: dict[str, str] = {}
    patron = re.compile(r"^([0-9a-fA-F]{64})  ([^/\\\r\n]+)$")
    for linea in lineas:
        coincidencia = patron.fullmatch(linea)
        if not coincidencia:
            raise ErrorRelease("El archivo SHA-256 tiene un formato inválido.")
        nombre = coincidencia.group(2)
        if nombre in sumas:
            raise ErrorRelease("El archivo SHA-256 contiene nombres duplicados.")
        sumas[nombre] = coincidencia.group(1).lower()
    return sumas


def _validar_envoltura_zip(ruta: Path) -> None:
    """Rechaza prefijos autoextraíbles y bytes anexos posteriores al EOCD."""

    tamano = ruta.stat().st_size
    if tamano < 22:
        raise ErrorRelease("El archivo ZIP es demasiado corto.")
    with ruta.open("rb") as archivo:
        if archivo.read(4) != b"PK\x03\x04":
            raise ErrorRelease("El ZIP contiene un prefijo no canónico.")
        longitud_cola = min(tamano, 22 + 65535)
        archivo.seek(tamano - longitud_cola)
        cola = archivo.read(longitud_cola)
    posicion = cola.rfind(b"PK\x05\x06")
    if posicion < 0 or posicion + 22 > len(cola):
        raise ErrorRelease("El ZIP no contiene un directorio final válido.")
    longitud_comentario = int.from_bytes(cola[posicion + 20:posicion + 22], "little")
    if posicion + 22 + longitud_comentario != len(cola):
        raise ErrorRelease("El ZIP contiene bytes anexos no canónicos.")


def _validar_manifiesto(
    manifiesto: object,
    nombre_zip: str,
) -> tuple[dict[str, object], dict[str, dict[str, object]]]:
    if not isinstance(manifiesto, dict):
        raise ErrorRelease("La raíz del manifiesto debe ser un objeto.")
    requeridas = {
        "artifact_name", "content_policy", "created_utc", "dependency_bundle",
        "file_count", "files",
        "format_version", "payload_size", "product", "release_version",
        "source_commit", "source_date_epoch", "source_dirty", "target",
    }
    if set(manifiesto) != requeridas:
        raise ErrorRelease("El manifiesto no contiene exactamente el esquema esperado.")
    if manifiesto["format_version"] != FORMATO_MANIFIESTO:
        raise ErrorRelease("Versión de formato de manifiesto no soportada.")
    if manifiesto["product"] != PRODUCTO:
        raise ErrorRelease("El producto del manifiesto no coincide.")
    if manifiesto["content_policy"] != POLITICA_CONTENIDO:
        raise ErrorRelease("La política de contenido del manifiesto no coincide.")
    if manifiesto["target"] != DESTINO_RELEASE:
        raise ErrorRelease("El destino de la release no es Windows CPython 3.13 x64.")
    if (
        not isinstance(manifiesto["dependency_bundle"], str)
        or manifiesto["dependency_bundle"] not in {"none", "wheelhouse"}
    ):
        raise ErrorRelease("El tipo de paquete de dependencias no es válido.")
    if manifiesto["artifact_name"] != nombre_zip:
        raise ErrorRelease("El nombre del ZIP no coincide con el manifiesto.")
    version = manifiesto["release_version"]
    if not isinstance(version, str):
        raise ErrorRelease("La versión del manifiesto no es textual.")
    _validar_version(version)
    if nombre_zip != f"{PRODUCTO}-{version}.zip":
        raise ErrorRelease("El nombre del artefacto no coincide con su versión.")
    commit = manifiesto["source_commit"]
    if not isinstance(commit, str) or not PATRON_COMMIT.fullmatch(commit):
        raise ErrorRelease("Commit inválido en el manifiesto.")
    epoch = manifiesto["source_date_epoch"]
    if (
        isinstance(epoch, bool) or not isinstance(epoch, int)
        or epoch < 315532800 or epoch > 4354819199
    ):
        raise ErrorRelease("Fecha reproducible inválida en el manifiesto.")
    if manifiesto["created_utc"] != _fecha_iso(epoch):
        raise ErrorRelease("La fecha legible no coincide con SOURCE_DATE_EPOCH.")
    if not isinstance(manifiesto["source_dirty"], bool):
        raise ErrorRelease("El indicador source_dirty no es booleano.")
    archivos = manifiesto["files"]
    if not isinstance(archivos, list) or not archivos:
        raise ErrorRelease("El manifiesto no declara archivos.")
    conteo = manifiesto["file_count"]
    if isinstance(conteo, bool) or not isinstance(conteo, int) or conteo != len(archivos):
        raise ErrorRelease("El conteo de archivos del manifiesto no coincide.")

    por_ruta: dict[str, dict[str, object]] = {}
    orden: list[bytes] = []
    tamano_total = 0
    for item in archivos:
        if not isinstance(item, dict) or set(item) != {"path", "sha256", "size"}:
            raise ErrorRelease("Entrada de archivo inválida en el manifiesto.")
        ruta = item["path"]
        sha = item["sha256"]
        tamano = item["size"]
        if not isinstance(ruta, str):
            raise ErrorRelease("Ruta no textual en el manifiesto.")
        ruta = _validar_ruta_relativa(ruta)
        if ruta == ENTRADA_MANIFIESTO:
            raise ErrorRelease("El manifiesto no puede declararse como payload.")
        partes = PurePosixPath(ruta).parts
        if (
            (
                any(_es_directorio_excluido(parte) for parte in partes[:-1])
                and ruta not in ARCHIVOS_EN_DIRECTORIOS_RESERVADOS
            )
            or _es_archivo_excluido(partes[-1])
        ):
            raise ErrorRelease(f"El manifiesto declara una ruta excluida: {ruta}.")
        clave = ruta.casefold()
        if clave in por_ruta:
            raise ErrorRelease("El manifiesto contiene rutas duplicadas o ambiguas.")
        if not isinstance(sha, str) or not PATRON_SHA256.fullmatch(sha):
            raise ErrorRelease(f"SHA-256 inválido para {ruta}.")
        if isinstance(tamano, bool) or not isinstance(tamano, int) or tamano < 0:
            raise ErrorRelease(f"Tamaño inválido para {ruta}.")
        tamano_total += tamano
        if tamano > TAMANO_MAXIMO_ARCHIVO or tamano_total > TAMANO_MAXIMO_PAYLOAD:
            raise ErrorRelease("El manifiesto excede los límites de tamaño.")
        por_ruta[clave] = item
        orden.append(ruta.encode("utf-8"))
    if orden != sorted(orden):
        raise ErrorRelease("Los archivos del manifiesto no están en orden canónico.")
    _validar_contrato_payload({str(item["path"]) for item in archivos})
    wheels: list[str] = []
    for item in archivos:
        ruta = str(item["path"])
        partes = PurePosixPath(ruta).parts
        if partes[0].casefold() != "wheelhouse":
            continue
        if (
            partes[0] != "wheelhouse"
            or len(partes) != 2
            or not partes[1].casefold().endswith(".whl")
        ):
            raise ErrorRelease(
                "El wheelhouse sólo admite wheels en un directorio plano canónico."
            )
        wheels.append(ruta)
    if bool(wheels) != (manifiesto["dependency_bundle"] == "wheelhouse"):
        raise ErrorRelease("La declaración del wheelhouse no coincide con el contenido.")
    payload_size = manifiesto["payload_size"]
    if (
        isinstance(payload_size, bool) or not isinstance(payload_size, int)
        or payload_size != tamano_total
    ):
        raise ErrorRelease("El tamaño total del manifiesto no coincide.")
    return manifiesto, por_ruta


def verificar_release(
    *,
    archivo_zip: Path,
    archivo_manifiesto: Path | None = None,
    archivo_sumas: Path | None = None,
) -> dict[str, object]:
    """Verifica hashes externos, manifiesto embebido y archivos sin extraer."""

    archivo_zip_solicitado = Path(archivo_zip)
    _rechazar_enlaces_en_componentes(archivo_zip_solicitado, "El ZIP")
    archivo_zip = archivo_zip_solicitado.resolve(strict=True)
    archivo_manifiesto = (
        Path(archivo_manifiesto)
        if archivo_manifiesto
        else archivo_zip.with_suffix(".manifest.json")
    )
    archivo_sumas = (
        Path(archivo_sumas)
        if archivo_sumas
        else archivo_zip.with_suffix(".sha256")
    )
    for solicitada in (archivo_manifiesto, archivo_sumas):
        _rechazar_enlaces_en_componentes(solicitada, "Los artefactos")
    archivo_manifiesto = archivo_manifiesto.resolve(strict=True)
    archivo_sumas = archivo_sumas.resolve(strict=True)
    for ruta in (archivo_zip, archivo_manifiesto, archivo_sumas):
        if _es_enlace(ruta) or not ruta.is_file():
            raise ErrorRelease("Los artefactos deben ser archivos regulares, no enlaces.")

    sumas = _leer_sumas(archivo_sumas)
    nombres = {archivo_zip.name, archivo_manifiesto.name}
    if set(sumas) != nombres:
        raise ErrorRelease("El archivo SHA-256 no declara exactamente ZIP y manifiesto.")
    if _sha256_archivo(archivo_zip) != sumas[archivo_zip.name]:
        raise ErrorRelease("El SHA-256 del ZIP no coincide.")
    if _sha256_archivo(archivo_manifiesto) != sumas[archivo_manifiesto.name]:
        raise ErrorRelease("El SHA-256 del manifiesto no coincide.")
    _validar_envoltura_zip(archivo_zip)
    if archivo_manifiesto.stat().st_size > TAMANO_MAXIMO_MANIFIESTO:
        raise ErrorRelease("El manifiesto excede el límite permitido.")

    bytes_manifiesto = archivo_manifiesto.read_bytes()
    manifiesto, por_ruta = _validar_manifiesto(
        _json_sin_duplicados(bytes_manifiesto), archivo_zip.name
    )
    fecha_esperada = _fecha_zip(int(manifiesto["source_date_epoch"]))
    esperadas = {item["path"] for item in manifiesto["files"]}
    esperadas.add(ENTRADA_MANIFIESTO)
    bytes_catalogo: bytes | None = None
    temporal_verificacion = tempfile.TemporaryDirectory(prefix="tocayos-wheel-verify-")
    raiz_temporal = Path(temporal_verificacion.name)
    wheelhouse_temporal = raiz_temporal / "wheelhouse"
    wheelhouse_temporal.mkdir()
    requirements_temporal: Path | None = None
    try:
        with zipfile.ZipFile(archivo_zip, "r") as paquete:
            if paquete.comment:
                raise ErrorRelease("El ZIP contiene un comentario no canónico.")
            infos = paquete.infolist()
            rutas_zip = [info.filename for info in infos]
            if len(rutas_zip) != len(set(rutas_zip)):
                raise ErrorRelease("El ZIP contiene nombres duplicados.")
            if set(rutas_zip) != esperadas:
                raise ErrorRelease("El contenido del ZIP no coincide con el manifiesto.")
            orden_esperado = [
                str(item["path"]) for item in manifiesto["files"]
            ] + [ENTRADA_MANIFIESTO]
            if rutas_zip != orden_esperado:
                raise ErrorRelease("Las entradas del ZIP no están en orden canónico.")
            for info in infos:
                ruta = _validar_ruta_relativa(info.filename)
                if info.is_dir() or info.flag_bits & 0x1:
                    raise ErrorRelease(f"Entrada ZIP no permitida: {ruta}.")
                tipo = (info.external_attr >> 16) & 0o170000
                if tipo not in {0, stat.S_IFREG}:
                    raise ErrorRelease(f"Entrada ZIP no regular: {ruta}.")
                if info.compress_type != zipfile.ZIP_DEFLATED:
                    raise ErrorRelease(f"Compresión no canónica para {ruta}.")
                if info.date_time != fecha_esperada:
                    raise ErrorRelease(f"Fecha ZIP no reproducible para {ruta}.")
                info_esperada = _info_zip(ruta, fecha_esperada)
                if (
                    info.create_system != 3
                    or info.external_attr != info_esperada.external_attr
                    or info.extra
                    or info.comment
                    or info.flag_bits & ~0x800
                ):
                    raise ErrorRelease(f"Metadatos ZIP no canónicos para {ruta}.")
                if ruta == ENTRADA_MANIFIESTO:
                    if info.file_size > TAMANO_MAXIMO_MANIFIESTO:
                        raise ErrorRelease("El manifiesto embebido excede el límite.")
                    if paquete.read(info) != bytes_manifiesto:
                        raise ErrorRelease("El manifiesto embebido no coincide con el externo.")
                    continue
                esperado = por_ruta[ruta.casefold()]
                if info.file_size != esperado["size"]:
                    raise ErrorRelease(f"El tamaño no coincide para {ruta}.")
                digest = hashlib.sha256()
                leidos = 0
                captura = None
                if ruta.startswith("wheelhouse/"):
                    captura = (wheelhouse_temporal / PurePosixPath(ruta).name).open("xb")
                elif ruta == "requirements-lock.txt":
                    requirements_temporal = raiz_temporal / "requirements-lock.txt"
                    captura = requirements_temporal.open("xb")
                try:
                    with paquete.open(info, "r") as contenido:
                        for bloque in iter(lambda: contenido.read(1024 * 1024), b""):
                            digest.update(bloque)
                            leidos += len(bloque)
                            if leidos > TAMANO_MAXIMO_ARCHIVO:
                                raise ErrorRelease(f"Entrada ZIP demasiado grande: {ruta}.")
                            if captura is not None:
                                captura.write(bloque)
                finally:
                    if captura is not None:
                        captura.close()
                if leidos != esperado["size"] or digest.hexdigest() != esperado["sha256"]:
                    raise ErrorRelease(f"El contenido no coincide para {ruta}.")
                if ruta == RUTA_CATALOGO_HISTORICO:
                    bytes_catalogo = paquete.read(info)
            if bytes_catalogo is None:
                raise ErrorRelease("El ZIP no contiene el catálogo histórico obligatorio.")
            _validar_catalogo_historico(
                io.BytesIO(bytes_catalogo), "El catálogo histórico del ZIP"
            )
            version_item = por_ruta["version"]
            try:
                version_payload = paquete.read(str(version_item["path"])).decode("ascii")
            except (KeyError, UnicodeDecodeError) as exc:
                raise ErrorRelease("VERSION dentro del ZIP no es válido.") from exc
            version_esperada = str(manifiesto["release_version"])
            if version_payload not in {
                version_esperada,
                version_esperada + "\n",
                version_esperada + "\r\n",
            }:
                raise ErrorRelease("VERSION dentro del ZIP no coincide con el manifiesto.")
        if requirements_temporal is None:
            raise ErrorRelease("El ZIP no contiene requirements-lock.txt verificable.")
        _validar_lock_exacto(requirements_temporal)
        if manifiesto["dependency_bundle"] == "wheelhouse":
            wheels_temporales = sorted(wheelhouse_temporal.iterdir())
            if not wheels_temporales:
                raise ErrorRelease("El wheelhouse declarado está vacío.")
            for wheel in wheels_temporales:
                _validar_wheel(wheel, f"El wheel {wheel.name} del ZIP")
            _validar_wheelhouse_exacto(wheelhouse_temporal, requirements_temporal)
    except ErrorRelease:
        raise
    except (zipfile.BadZipFile, RuntimeError) as exc:
        raise ErrorRelease("El archivo ZIP está dañado o no es compatible.") from exc
    finally:
        temporal_verificacion.cleanup()
    return manifiesto


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Construye o verifica una release reproducible del servidor Edge."
    )
    comandos = parser.add_subparsers(dest="comando", required=True)
    construir = comandos.add_parser("build", help="Construir una release.")
    construir.add_argument("--version", required=True)
    construir.add_argument("--source", type=Path, default=Path.cwd())
    construir.add_argument("--output", type=Path, default=Path("release") / "servidor")
    construir.add_argument("--commit")
    construir.add_argument("--source-date-epoch", type=int)
    construir.add_argument(
        "--include", action="append", default=[], metavar="RUTA",
        help="Añadir una ruta relativa a la lista permitida predeterminada.",
    )
    construir.add_argument("--allow-dirty", action="store_true")
    construir.add_argument("--force", action="store_true")
    dependencias = construir.add_mutually_exclusive_group(required=True)
    dependencias.add_argument(
        "--wheelhouse",
        type=Path,
        help="Directorio plano de wheels ya descargados para instalación offline.",
    )
    dependencias.add_argument(
        "--source-only",
        action="store_true",
        help="Crear deliberadamente un paquete sin dependencias (no sirve para instalación limpia).",
    )
    verificar = comandos.add_parser("verify", help="Verificar una release sin extraerla.")
    verificar.add_argument("--archive", type=Path, required=True)
    verificar.add_argument("--manifest", type=Path)
    verificar.add_argument("--checksum", type=Path)
    return parser


def main(argumentos: list[str] | None = None) -> int:
    opciones = _parser().parse_args(argumentos)
    try:
        if opciones.comando == "build":
            origen = opciones.source
            rutas = None
            if opciones.include:
                faltantes = [
                    ruta for ruta in RUTAS_REQUERIDAS if not (origen / ruta).exists()
                ]
                if faltantes:
                    raise ErrorRelease("Faltan rutas requeridas: " + ", ".join(faltantes))
                rutas = RUTAS_REQUERIDAS + tuple(
                    ruta for ruta in RUTAS_OPCIONALES if (origen / ruta).exists()
                ) + tuple(opciones.include)
            resultado = crear_release(
                origen=origen,
                destino=opciones.output,
                version=opciones.version,
                commit=opciones.commit,
                source_date_epoch=opciones.source_date_epoch,
                permitir_sucio=opciones.allow_dirty,
                sobrescribir=opciones.force,
                rutas_incluidas=rutas,
                wheelhouse=opciones.wheelhouse,
            )
            print(json.dumps({
                "manifest": str(resultado.manifiesto),
                "manifest_sha256": resultado.sha256_manifiesto,
                "sha256_file": str(resultado.sumas),
                "zip": str(resultado.zip),
                "zip_sha256": resultado.sha256_zip,
            }, ensure_ascii=False, indent=2, sort_keys=True))
        else:
            manifiesto = verificar_release(
                archivo_zip=opciones.archive,
                archivo_manifiesto=opciones.manifest,
                archivo_sumas=opciones.checksum,
            )
            print(json.dumps({
                "commit": manifiesto["source_commit"],
                "files": manifiesto["file_count"],
                "status": "ok",
                "version": manifiesto["release_version"],
            }, ensure_ascii=False, sort_keys=True))
        return 0
    except (ErrorRelease, FileNotFoundError, OSError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
