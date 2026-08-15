from pathlib import Path
import warnings

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from catalogo.models import Categoria, Producto
from personas.models import Sucursal


class Command(BaseCommand):
    help = "Importa los nombres confiables del XLSX legado como productos inactivos para revisión."

    def add_arguments(self, parser):
        parser.add_argument("archivo", type=Path)
        parser.add_argument("--sucursal", default="ARBOLEDAS")

    def handle(self, *args, **options):
        archivo = options["archivo"]
        if not archivo.is_file():
            raise CommandError(f"No existe el archivo: {archivo}")
        try:
            sucursal = Sucursal.objects.get(clave=options["sucursal"])
        except Sucursal.DoesNotExist as exc:
            raise CommandError("Primero ejecuta cargar_datos_iniciales.") from exc

        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            libro = load_workbook(archivo, read_only=True, data_only=True)
        hoja = libro["Productos"]
        categoria_revision, _ = Categoria.objects.get_or_create(
            sucursal=sucursal,
            nombre="Por revisar (legado)",
            defaults={"orden": 99, "activa": False},
        )
        creados = 0
        for indice, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), 1):
            codigo, nombre, _, _, destino = (list(fila) + [None] * 5)[:5]
            if not nombre:
                continue
            codigo_base = str(codigo or f"LEG-{indice:03d}").strip()[:30]
            if Producto.objects.filter(sucursal=sucursal, nombre__iexact=str(nombre).strip()).exists():
                continue
            codigo_final = codigo_base
            sufijo = 1
            while Producto.objects.filter(sucursal=sucursal, codigo=codigo_final).exists():
                sufijo += 1
                codigo_final = f"{codigo_base[:25]}-{sufijo}"
            Producto.objects.create(
                sucursal=sucursal,
                categoria=categoria_revision,
                codigo=codigo_final,
                nombre=str(nombre).strip(),
                nombre_corto=codigo_final[:24],
                destino_impresion="barra" if str(destino).strip().lower() == "barra" else "cocina",
                activo=False,
                origen="legado_incompleto",
            )
            creados += 1
        self.stdout.write(self.style.SUCCESS(f"{creados} nombres agregados como inactivos; no se importaron precios."))
