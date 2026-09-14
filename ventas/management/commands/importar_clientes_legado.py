import re
import warnings
from collections import OrderedDict
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from personas.models import Sucursal
from ventas.clientes import (
    MAX_DOMICILIOS_POR_CLIENTE,
    MAX_TELEFONOS_POR_CLIENTE,
    cliente_payload,
    guardar_cliente,
)
from ventas.models import Cliente
from ventas.normalizacion import normalizar_texto, normalizar_telefono


MARCADOR_ORIGEN = "Importado desde Clientes-Domicilio-Completo.xlsx"
PATRON_CONTACTO = re.compile(r"(?<!\d)(33\d{8}|\d{8})(?!\d)")
PATRON_INTERIOR = re.compile(
    r"\b(INT(?:ERIOR)?|DEP(?:ARTAMENTO)?|DEPTO|CASA|LOCAL)\.?\s*#?\s*([A-Z0-9-]+)",
    re.IGNORECASE,
)


def texto(valor):
    return str(valor or "").replace("�", "Ñ").strip()


def separar_domicilio(valor):
    domicilio = re.sub(r"\s+", " ", texto(valor))
    interior = ""
    coincidencia_interior = PATRON_INTERIOR.search(domicilio)
    base = domicilio
    if coincidencia_interior:
        interior = f"{coincidencia_interior.group(1).upper()} {coincidencia_interior.group(2)}"
        base = f"{domicilio[:coincidencia_interior.start()]} {domicilio[coincidencia_interior.end():]}".strip()
    numeros = list(re.finditer(r"(?<![A-Z0-9])#?(\d+[A-Z]?)(?![A-Z0-9])", base, re.IGNORECASE))
    if not numeros:
        return domicilio, "", interior
    exterior = numeros[-1]
    calle = f"{base[:exterior.start()]} {base[exterior.end():]}".strip(" ,-#")
    calle = re.sub(r"\s+", " ", calle) or domicilio
    return calle, exterior.group(1), interior


def clave_domicilio(datos):
    return normalizar_texto(
        " ".join(
            str(datos.get(campo, ""))
            for campo in (
                "calle",
                "numero_exterior",
                "numero_interior",
                "colonia",
                "codigo_postal",
                "municipio",
                "referencia",
            )
        )
    )


def notas_con_origen(notas):
    notas = str(notas or "").strip()
    if MARCADOR_ORIGEN in notas:
        return notas
    return "\n".join(parte for parte in (notas, MARCADOR_ORIGEN) if parte)


class Command(BaseCommand):
    help = "Importa clientes completos del directorio XLSX legado, agrupados por nombre."

    def add_arguments(self, parser):
        parser.add_argument("archivo", type=Path)
        parser.add_argument("--sucursal", default="ARBOLEDAS")
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument(
            "--actualizar-existentes",
            action="store_true",
            help="Fusiona datos del XLSX en clientes que ya coinciden por nombre.",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        archivo = options["archivo"]
        if not archivo.is_file():
            raise CommandError(f"No existe el archivo: {archivo}")
        try:
            sucursal = Sucursal.objects.get(clave=options["sucursal"])
        except Sucursal.DoesNotExist as exc:
            raise CommandError("Primero aprovisiona la sucursal solicitada.") from exc

        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            libro = load_workbook(archivo, read_only=True, data_only=True)
        if "Clientes" not in libro.sheetnames:
            raise CommandError("El archivo no contiene la hoja Clientes.")

        grupos = OrderedDict()
        omitidos = 0
        for fila in libro["Clientes"].iter_rows(min_row=2, values_only=True):
            nombre, domicilio, colonia, ciudad, _, telefono, referencia, *_ = (list(fila) + [None] * 9)[:9]
            nombre, domicilio = texto(nombre), texto(domicilio)
            telefono = texto(telefono)
            referencia = texto(referencia)
            telefono_normalizado = normalizar_telefono(telefono)
            contacto_rotativo = PATRON_CONTACTO.search(referencia) if len(telefono_normalizado) < 7 else None
            if not nombre:
                omitidos += 1
                continue
            clave = normalizar_texto(nombre)
            grupo = grupos.setdefault(
                clave,
                {
                    "nombre": nombre,
                    "telefonos": OrderedDict(),
                    "domicilios": OrderedDict(),
                    "comentarios_multiples": False,
                },
            )
            if len(telefono_normalizado) >= 7:
                grupo["telefonos"].setdefault(
                    telefono_normalizado,
                    {"etiqueta": "Principal", "numero": telefono},
                )
            if contacto_rotativo:
                grupo["comentarios_multiples"] = True
            if domicilio:
                calle, exterior, interior = separar_domicilio(domicilio)
                domicilio_clave = normalizar_texto(" ".join([domicilio, texto(colonia), texto(ciudad)]))
                grupo["domicilios"].setdefault(
                    domicilio_clave,
                    {
                        "etiqueta": "Principal",
                        "calle": calle,
                        "numero_exterior": exterior,
                        "numero_interior": interior,
                        "colonia": texto(colonia),
                        "codigo_postal": "",
                        "municipio": texto(ciudad),
                        "referencia": referencia,
                    },
                )

        libro.close()

        creados = actualizados = existentes = 0
        for grupo in grupos.values():
            candidatos = list(
                Cliente.objects.filter(
                    sucursal=sucursal,
                    nombre_normalizado=normalizar_texto(grupo["nombre"]),
                    activo=True,
                )
                .prefetch_related("telefonos", "domicilios")
                .order_by("-actualizado_en")[:10]
            )
            cliente = next(
                (item for item in candidatos if MARCADOR_ORIGEN in item.notas),
                candidatos[0] if candidatos else None,
            )
            if cliente and not options["actualizar_existentes"]:
                existentes += 1
                continue
            telefonos = OrderedDict()
            domicilios = OrderedDict()
            notas = MARCADOR_ORIGEN
            comentarios_multiples = grupo["comentarios_multiples"]
            if cliente:
                existente = cliente_payload(cliente)
                notas = notas_con_origen(existente["notas"])
                comentarios_multiples = (
                    existente["comentarios_multiples"] or comentarios_multiples
                )
                for telefono in existente["telefonos"]:
                    telefonos[normalizar_telefono(telefono["numero"])] = telefono
                for domicilio in existente["domicilios"]:
                    domicilios[clave_domicilio(domicilio)] = domicilio
            for telefono in grupo["telefonos"].values():
                telefonos.setdefault(
                    normalizar_telefono(telefono["numero"]),
                    telefono,
                )
            for domicilio in grupo["domicilios"].values():
                domicilios.setdefault(clave_domicilio(domicilio), domicilio)

            datos = {
                "nombre": grupo["nombre"],
                "notas": notas,
                "comentarios_multiples": comentarios_multiples,
                "telefonos": list(telefonos.values())[:MAX_TELEFONOS_POR_CLIENTE],
                "domicilios": list(domicilios.values())[:MAX_DOMICILIOS_POR_CLIENTE],
            }
            guardar_cliente(sucursal, datos, cliente=cliente)
            if cliente:
                actualizados += 1
            else:
                creados += 1

        if options["dry_run"]:
            transaction.set_rollback(True)
        modo = "simulados" if options["dry_run"] else "importados"
        self.stdout.write(
            self.style.SUCCESS(
                f"{creados} clientes nuevos y {actualizados} actualizados {modo}; "
                f"{existentes} ya existentes sin cambios; "
                f"{omitidos} filas incompletas omitidas."
            )
        )
